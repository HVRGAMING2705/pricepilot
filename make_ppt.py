from pptx import Presentation
from pptx.util import Inches, Pt
prs = Presentation(); prs.slide_width = Inches(13.333); prs.slide_height = Inches(7.5)
BLANK = prs.slide_layouts[6]
def slide(title, bullets, sub=None):
    s = prs.slides.add_slide(BLANK)
    t = s.shapes.add_textbox(Inches(0.6), Inches(0.4), Inches(12), Inches(1))
    t.text_frame.word_wrap = True; t.text_frame.text = title
    t.text_frame.paragraphs[0].font.size = Pt(36); t.text_frame.paragraphs[0].font.bold = True
    y = 1.6
    if sub:
        b = s.shapes.add_textbox(Inches(0.6), Inches(y), Inches(12), Inches(0.7)); b.text_frame.word_wrap=True; b.text_frame.text=sub; b.text_frame.paragraphs[0].font.size=Pt(20); y += 0.9
    for line in bullets:
        b = s.shapes.add_textbox(Inches(0.8), Inches(y), Inches(11.8), Inches(0.8)); b.text_frame.word_wrap=True; b.text_frame.text="•  "+line; b.text_frame.paragraphs[0].font.size=Pt(20); y += 0.85
    return s
slide("PricePilot", ["Used Car Price Predictor"], "An AI/ML project by Vikranth Reddy — B.Tech CSE (AI & ML), Woxsen University")
slide("The Problem", ["Selling a used car starts with guesswork — owners overprice and wait, or underprice and lose money", "Buyers have no quick way to check whether an asking price is fair", "Online listings vary wildly for cars that look almost identical"])
slide("The Solution", ["PricePilot predicts a used car's resale price (in ₹ lakh) from 7 inputs", "Age, mileage, engine size, brand tier, owners, fuel, transmission", "Interactive demo: change any input, see the predicted price instantly — fully in the browser"])
slide("Tech Stack", ["Python + scikit-learn: ridge regression (alpha = 1.0)", "Synthetic dataset: 600 car listings (seed 42)", "HTML + CSS + vanilla JavaScript demo", "GitHub + Vercel hosting"])
slide("Architecture", ["Inputs: 7 car features → standardization (train mean / std) → ridge regression → clip to ₹0.50–80 lakh", "Model exported as model.json; the browser demo reproduces it exactly", "Strongest drivers: age (depreciation) and brand tier"])
slide("Demo", ["Try the presets: almost-new premium, average family car, old budget car", "Almost-new premium (1 yr, 10,000 km) → ₹14.27 lakh", "Average family car (5 yrs, 60,000 km) → ₹6.30 lakh"])
slide("Results", ["5-fold cross-validated R² = 0.745", "Mean absolute error = ₹1.79 lakh (baseline: ₹3.62 lakh — about 2× better)", "RMSE = ₹2.30 lakh", "Strongest drivers: car age and brand tier, as expected in real markets"])
slide("Future Scope", ["Train on real listings with condition and service-history features", "Add city-wise price adjustments and price-trend charts", "Explain each prediction: which feature raised or lowered the price", "Photo-based condition scoring from car images"])
slide("Thank You", ["PricePilot — know the fair price before you buy or sell."], "Vikranth Reddy · github.com/HVRGAMING2705/pricepilot")
out = "/home/hatch/workspace/goals/daily-ai-ml-project-factory/files/AIML_PricePilot_2026-10-26.pptx"
prs.save(out); print("saved", out, "slides:", len(prs.slides._sldIdLst))
